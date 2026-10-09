"""Construction des exports : aplatis les lignes, sérialise en JSON/CSV/XLSX,
écrit le fichier sous `settings.EXPORT_DIR` et crée la ligne `exports` (statut ready).

`build_export(session, project_id, kind, format, filters) -> (export_id, file_bytes)`.
"""

from __future__ import annotations

import csv
import io
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from webintel.config.settings import get_settings
from webintel.core.logging import get_logger
from webintel.core.utils import utcnow
from webintel.database.models import Alert, CrawlJob, Entity, Event, Export, Monitor, PageSnapshot, SourceObservation
from webintel.search.engine import search_entities

logger = get_logger(__name__)

_EXTENSIONS = {"json": "json", "csv": "csv", "xlsx": "xlsx"}


# ---------------------------------------------------------------------------
# Aplatissement & sérialisation
# ---------------------------------------------------------------------------

def flatten_dict(d: dict, prefix: str = "") -> dict:
    """Aplatit un dict imbriqué : `price.amount` -> clé `price_amount`.

    Les listes sont jointes avec `"; "` (éléments dict -> JSON compact),
    les dicts sont récursés, les scalaires sont conservés tels quels.
    """
    flat: dict = {}
    for key, value in (d or {}).items():
        name = f"{prefix}_{key}" if prefix else str(key)
        if isinstance(value, dict):
            flat.update(flatten_dict(value, name))
        elif isinstance(value, (list, tuple)):
            parts = []
            for item in value:
                if isinstance(item, dict):
                    parts.append(json.dumps(item, ensure_ascii=False, default=str))
                elif item is None:
                    continue
                else:
                    parts.append(str(item))
            flat[name] = "; ".join(parts)
        elif isinstance(value, datetime):
            flat[name] = value.isoformat()
        else:
            flat[name] = value
    return flat


def to_json(rows: list[dict]) -> bytes:
    """Sérialise les lignes en JSON UTF-8 indenté (sans échappement ASCII)."""
    return json.dumps(rows, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def to_csv(rows: list[dict]) -> bytes:
    """Sérialise les lignes en CSV UTF-8 avec BOM (lisible par Excel)."""
    buffer = io.StringIO()
    headers = _union_headers(rows)
    writer = csv.DictWriter(buffer, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({h: _cell(row.get(h)) for h in headers})
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(rows: list[dict]) -> bytes:
    """Sérialise les lignes en XLSX (en-têtes en gras, largeurs auto simples)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "export"
    headers = _union_headers(rows)
    bold = Font(bold=True)
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = bold
    for row_idx, row in enumerate(rows, start=2):
        for col, header in enumerate(headers, start=1):
            sheet.cell(row=row_idx, column=col, value=_cell(row.get(header)))
    # Largeurs de colonnes : max(en-tête, échantillon) borné.
    for col, header in enumerate(headers, start=1):
        width = len(str(header))
        for row in rows[:50]:
            width = max(width, len(str(row.get(header) or "")))
        sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = min(width + 2, 60)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _union_headers(rows: list[dict]) -> list[str]:
    """Union ordonnée des clés de toutes les lignes (stabilité du CSV/XLSX)."""
    headers: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)
    return headers or ["(vide)"]


def _cell(value):
    """Valeur de cellule CSV/XLSX : scalaires et JSON compact pour le reste."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, (str, int, float)):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)


# ---------------------------------------------------------------------------
# Collecte des lignes selon le type d'export
# ---------------------------------------------------------------------------

def _rows_entities(session: Session, project_id: int, filters: dict) -> list[dict]:
    entities = search_entities(
        session,
        project_id=project_id,
        q=filters.get("q"),
        entity_type=filters.get("type"),
        min_price=filters.get("min_price"),
        max_price=filters.get("max_price"),
        domain=filters.get("domain"),
        limit=int(filters.get("limit", 1000)),
    )
    # Enrichit avec les données complètes de l'entité (aplaties).
    rows: list[dict] = []
    for item in entities:
        entity = session.get(Entity, item["id"])
        base = {k: v for k, v in item.items() if k not in ("domains",)}
        base.update(flatten_dict({"donnees": entity.data or {}}))
        base["sources"] = "; ".join(item.get("sources") or [])
        rows.append(base)
    return rows


def _rows_events(session: Session, project_id: int, filters: dict) -> list[dict]:
    query = session.query(Event).filter(Event.project_id == project_id)
    if filters.get("type"):
        query = query.filter(Event.type == str(filters["type"]).lower())
    events = query.order_by(Event.created_at.desc()).limit(int(filters.get("limit", 1000))).all()
    return [flatten_dict({
        "id": e.id, "type": e.type, "entity_id": e.entity_id, "url": e.url,
        "importance": e.importance, "before": e.before, "after": e.after,
        "evidence": e.evidence, "created_at": e.created_at,
    }) for e in events]


def _rows_snapshots(session: Session, project_id: int, filters: dict) -> list[dict]:
    query = session.query(PageSnapshot).filter(PageSnapshot.project_id == project_id)
    if filters.get("url"):
        query = query.filter(PageSnapshot.url.contains(str(filters["url"])))
    snaps = query.order_by(PageSnapshot.fetched_at.desc()).limit(int(filters.get("limit", 1000))).all()
    rows = []
    for snap in snaps:
        payload = snap.payload or {}
        row = {
            "id": snap.id, "url": snap.url, "canonical_url": snap.canonical_url,
            "fetched_at": snap.fetched_at, "status_code": snap.status_code,
            "title": snap.title, "page_type": snap.page_type,
            "page_type_score": snap.page_type_score, "word_count": snap.word_count,
            "seo_score": (payload.get("seo") or {}).get("score"),
            "health_overall": ((payload.get("health") or {}).get("overall")),
            "technologies": "; ".join(
                str(t.get("technology")) for t in (payload.get("technologies") or [])
                if isinstance(t, dict) and t.get("technology")
            ),
        }
        rows.append(row)
    return rows


def _rows_monitors(session: Session, project_id: int, filters: dict) -> list[dict]:
    monitors = (
        session.query(Monitor)
        .filter(Monitor.project_id == project_id)
        .order_by(Monitor.id)
        .all()
    )
    return [flatten_dict({
        "id": m.id, "name": m.name, "target_type": m.target_type, "target": m.target,
        "interval_minutes": m.interval_minutes, "active": m.active,
        "last_run_at": m.last_run_at, "config": m.config, "created_at": m.created_at,
    }) for m in monitors]


def _rows_alerts(session: Session, project_id: int, filters: dict) -> list[dict]:
    alerts = (
        session.query(Alert)
        .filter(Alert.project_id == project_id)
        .order_by(Alert.created_at.desc())
        .limit(int(filters.get("limit", 1000)))
        .all()
    )
    return [flatten_dict({
        "id": a.id, "monitor_id": a.monitor_id, "channel": a.channel,
        "status": a.status, "rule": a.rule, "payload": a.payload,
        "error": a.error, "created_at": a.created_at,
    }) for a in alerts]


def _rows_observations(session: Session, project_id: int, filters: dict) -> list[dict]:
    """Export des observations par source (P0-3) : une ligne par observation.

    Filtres optionnels : `field` (ex. "price"), `source_domain`.
    """
    query = (
        session.query(SourceObservation, Entity)
        .join(Entity, SourceObservation.entity_id == Entity.id)
        .filter(Entity.project_id == project_id)
    )
    if filters.get("field"):
        query = query.filter(SourceObservation.field == str(filters["field"]))
    if filters.get("source_domain"):
        query = query.filter(SourceObservation.source_domain == str(filters["source_domain"]))
    rows = []
    for obs, entity in query.order_by(
        SourceObservation.observed_at.desc(), SourceObservation.id.desc()
    ).limit(int(filters.get("limit", 5000))).all():
        rows.append(flatten_dict({
            "id": obs.id,
            "entity_id": obs.entity_id,
            "entity_key": entity.entity_key,
            "entity_type": entity.entity_type,
            "entity_name": entity.name,
            "source_id": obs.source_id,
            "source_url": obs.source_url,
            "source_domain": obs.source_domain,
            "field": obs.field,
            "value": obs.value,
            "value_hash": obs.value_hash,
            "confidence": obs.confidence,
            "snapshot_id": obs.snapshot_id,
            "observed_at": obs.observed_at,
        }))
    return rows


def _rows_jobs(session: Session, project_id: int, filters: dict) -> list[dict]:
    jobs = (
        session.query(CrawlJob)
        .filter(CrawlJob.project_id == project_id)
        .order_by(CrawlJob.created_at.desc())
        .limit(int(filters.get("limit", 1000)))
        .all()
    )
    return [flatten_dict({
        "id": j.id, "kind": j.kind, "status": j.status, "params": j.params,
        "progress": j.progress, "error": j.error,
        "created_at": j.created_at, "updated_at": j.updated_at,
    }) for j in jobs]


_COLLECTORS = {
    "entities": _rows_entities,
    "events": _rows_events,
    "snapshots": _rows_snapshots,
    "monitors": _rows_monitors,
    "alerts": _rows_alerts,
    "jobs": _rows_jobs,
    "observations": _rows_observations,
}

_SERIALIZERS = {"json": to_json, "csv": to_csv, "xlsx": to_xlsx}


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

def build_export(
    session: Session,
    project_id: int,
    kind: str,
    format: str,
    filters: dict | None = None,
) -> tuple[int, bytes]:
    """Construit un export et retourne `(export_id, file_bytes)`.

    - récupère les lignes selon `kind` (entities via `search.engine`, requêtes
      simples sinon), les aplatie avec `flatten_dict` ;
    - sérialise selon `format` (`json` / `csv` / `xlsx`) ;
    - écrit le fichier sous `settings.EXPORT_DIR` (dossier créé si besoin) ;
    - crée la ligne `exports` (statut `ready`).
    """
    kind = (kind or "").lower()
    format = (format or "").lower()
    if kind not in _COLLECTORS:
        raise ValueError(f"Type d'export inconnu : {kind!r} (attendus : {', '.join(sorted(_COLLECTORS))}).")
    if format not in _SERIALIZERS:
        raise ValueError(f"Format d'export inconnu : {format!r} (attendus : json, csv, xlsx).")
    filters = dict(filters or {})

    rows = _COLLECTORS[kind](session, project_id, filters)
    flat_rows = [flatten_dict(row) for row in rows]
    file_bytes = _SERIALIZERS[format](flat_rows)

    export_dir = get_settings().EXPORT_DIR
    os.makedirs(export_dir, exist_ok=True)
    stamp = utcnow().strftime("%Y%m%d-%H%M%S")
    filename = f"webintel-{kind}-{project_id}-{stamp}.{_EXTENSIONS[format]}"
    file_path = os.path.join(export_dir, filename)
    with open(file_path, "wb") as fh:
        fh.write(file_bytes)

    row = Export(
        project_id=project_id,
        kind=kind,
        format=format,
        params=filters,
        status="ready",
        file_path=file_path,
    )
    session.add(row)
    session.flush()
    logger.info("Export %s construit : %s (%d lignes, %d octets).",
                row.id, filename, len(flat_rows), len(file_bytes))
    return row.id, file_bytes
